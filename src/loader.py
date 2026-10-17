import bisect
import io
from pathlib import Path

MAX_BYTES = 15 * 1024 * 1024
MIN_CHARS = 200


class UploadError(ValueError):
    pass


def page_for(offset, page_starts):
    if not page_starts:
        return None
    return bisect.bisect_right(page_starts, offset)


def _check_text(text):
    if len(text.strip()) < MIN_CHARS:
        raise UploadError("No readable text found. If this is a scanned PDF, OCR is not supported yet.")
    return text


def _pdf(data):
    import pdfplumber

    parts, starts, pos = [], [], 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                t = (page.extract_text() or "").replace("\x00", "")
                starts.append(pos)
                parts.append(t)
                pos += len(t) + 2
    except Exception as e:
        raise UploadError(f"Could not read this PDF: {e}")
    return _check_text("\n\n".join(parts)), starts


def _docx(data):
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise UploadError(f"Could not read this Word file: {e}")
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            paras.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _check_text("\n\n".join(paras).replace("\x00", "")), None


def _txt(data):
    return _check_text(data.decode("utf-8", errors="replace").replace("\x00", "")), None


def extract_text(filename, data):
    if len(data) > MAX_BYTES:
        raise UploadError("File is larger than 15 MB.")
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext == ".txt":
        return _txt(data)
    raise UploadError("Only PDF, DOCX and TXT files are supported.")